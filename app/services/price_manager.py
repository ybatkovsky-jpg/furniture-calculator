import re
import unicodedata
from pathlib import Path
from typing import Any, Dict, Optional

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.price import PriceHistory, PriceItem

CATEGORY_MARKERS = {
    "лдсп": "ЛДСП",
    "мдф": "МДФ",
    "кромк": "Кромка",
    "петл": "Петли",
    "фурнитур": "Фурнитура",
    "ящик": "Ящики",
    "ящики": "Ящики",
    "стекл": "Стекло",
    "столешн": "Столешницы",
    "фасад": "Фасады",
    "алюмин": "Алюминиевый фасад",
    "комплектующ": "Комплектующие",
    "ручк": "Ручки",
    "профил": "Профили",
}

DEFAULT_PRICE_UNIT = "шт"


def normalize_text(value: Optional[Any]) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    text = unicodedata.normalize("NFKD", text)
    text = text.casefold()
    text = re.sub(r"[\s\W]+", " ", text)
    return text.strip()


def parse_number(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    text = text.replace("\u00a0", " ").replace(" ", "").replace(",", ".")
    cleaned = re.sub(r"[^0-9.\-]", "", text)
    if not cleaned or cleaned in {".", "-", "-."}:
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def detect_category_from_text(value: Optional[Any]) -> Optional[str]:
    text = normalize_text(value)
    for marker, category in CATEGORY_MARKERS.items():
        if marker in text:
            return category
    return None


def detect_price_unit(name: Optional[Any], category: Optional[Any], header: Optional[str], unit_cell: Optional[Any]) -> str:
    header_text = normalize_text(header)
    category_text = normalize_text(category)
    name_text = normalize_text(name)
    unit_text = normalize_text(unit_cell)

    if "м2" in header_text or "м²" in str(header or "") or "кв" in header_text:
        return "м²"
    if "м2" in category_text or "м²" in str(category or "") or "кв" in category_text:
        return "м²"
    if "м2" in name_text or "м²" in str(name or "") or "кв" in name_text:
        return "м²"
    if "шт" in header_text or "шт" in category_text or "шт" in name_text or "шт" in unit_text:
        return "шт"
    if "м" in header_text and "м2" not in header_text:
        return "м"
    if "м" in category_text and "м2" not in category_text:
        return "м"
    if "м" in name_text and "м2" not in name_text:
        return "м"
    if "кром" in category_text or "кром" in name_text:
        return "м"
    if "стекл" in category_text or "стекл" in name_text or "фасад" in category_text:
        return "м²"
    return DEFAULT_PRICE_UNIT


def guess_header_row(sheet) -> tuple[int, dict[str, int]]:
    header_patterns = {
        "name": ["наимен", "назван", "товар", "позици"],
        "brand": ["бренд", "производ", "марка"],
        "category": ["категор", "раздел", "группа"],
        "price": ["цен", "р/м", "р/м2", "р/м²", "цена/м", "цена/м²", "цена"],
        "unit": ["ед", "шт", "м2", "м²", "м", "кв"],
    }

    for row_index in range(1, min(12, sheet.max_row) + 1):
        row = [normalize_text(cell.value) for cell in sheet[row_index]]
        if not any(row):
            continue

        mapping: dict[str, int] = {}
        for index, cell_text in enumerate(row):
            if not cell_text:
                continue
            for name, patterns in header_patterns.items():
                if any(pattern in cell_text for pattern in patterns):
                    if name not in mapping:
                        mapping[name] = index
        if "name" in mapping and "price" in mapping:
            return row_index, mapping

    # Fallback positions when headers are not obvious
    return 1, {"name": 0, "brand": 1, "category": 2, "price": 3, "unit": 4}


def is_section_header(row: list[Any]) -> bool:
    if not row or row[0] is None or not str(row[0]).strip():
        return False

    rest = row[1:]
    return all(
        v is None or (isinstance(v, str) and not v.strip()) or (isinstance(v, (int, float)) and v == 0)
        for v in rest
    )


async def import_price_from_xlsx(
    path: str,
    session: AsyncSession,
    changed_by: Optional[str] = None,
) -> Dict[str, int]:
    """Импорт прайса из XLSX в базу данных."""
    workbook = load_workbook(Path(path), data_only=True)
    created = 0
    updated = 0
    skipped = 0

    for sheet in workbook.worksheets:
        current_category = detect_category_from_text(sheet.title) or None
        header_row_index, header_mapping = guess_header_row(sheet)

        for row_index in range(header_row_index + 1, sheet.max_row + 1):
            row = [sheet.cell(row=row_index, column=col_index).value for col_index in range(1, max(sheet.max_column, 6) + 1)]
            if not any(row):
                continue

            if is_section_header(row):
                section_text = row[0]
                category_hint = detect_category_from_text(section_text)
                if category_hint:
                    current_category = category_hint
                continue

            name = row[header_mapping.get("name", 0)] if header_mapping.get("name") is not None else row[0]
            if name is None or not str(name).strip():
                skipped += 1
                continue

            name = str(name).strip()
            category_cell = row[header_mapping.get("category")] if header_mapping.get("category") is not None else None
            category_text = str(category_cell).strip() if category_cell is not None else ""
            category_hint = detect_category_from_text(category_text)
            if category_hint:
                category = category_hint
                subcategory = None
            else:
                category = current_category or category_text or "Прочее"
                subcategory = category_text if category_text and category_text != category else None

            brand_cell = row[header_mapping.get("brand")] if header_mapping.get("brand") is not None else None
            brand = str(brand_cell).strip() if brand_cell is not None and str(brand_cell).strip() else None
            if not brand:
                brand_guess = str(name).split(maxsplit=1)[0]
                if brand_guess.upper() in {"EGGER", "EXTRAVERT", "LAMARTY", "BOYARD", "BLUM", "HETTICH"}:
                    brand = brand_guess

            price_cell = row[header_mapping.get("price")] if header_mapping.get("price") is not None else None
            price = parse_number(price_cell)
            if price is None:
                for value in row:
                    parsed = parse_number(value)
                    if parsed is not None and parsed > 0:
                        price = parsed
                        break

            unit_cell = row[header_mapping.get("unit")] if header_mapping.get("unit") is not None else None
            price_unit = detect_price_unit(name, category, sheet.cell(row=header_row_index, column=header_mapping.get("price", 4)).value if header_mapping.get("price") is not None else None, unit_cell)

            requires_manual_price = 1 if price is None or price == 0 else 0
            unit_price = None if requires_manual_price else float(price)

            query = select(PriceItem).where(
                PriceItem.name == name,
                PriceItem.category == category,
                PriceItem.brand == brand,
            )
            result = await session.execute(query)
            price_item = result.scalars().first()

            if price_item is None:
                price_item = PriceItem(
                    name=name,
                    category=category,
                    subcategory=subcategory,
                    brand=brand,
                    unit_price=unit_price,
                    price_unit=price_unit,
                    requires_manual_price=requires_manual_price,
                )
                session.add(price_item)
                created += 1
                continue

            has_changed = False
            old_price = float(price_item.unit_price or 0)
            new_price = float(price) if unit_price is not None else 0.0

            if price_item.unit_price != unit_price:
                price_item.unit_price = unit_price
                has_changed = True
            if price_item.price_unit != price_unit:
                price_item.price_unit = price_unit
                has_changed = True
            if price_item.category != category:
                price_item.category = category
                has_changed = True
            if price_item.subcategory != subcategory:
                price_item.subcategory = subcategory
                has_changed = True
            if price_item.brand != brand:
                price_item.brand = brand
                has_changed = True
            if price_item.requires_manual_price != requires_manual_price:
                price_item.requires_manual_price = requires_manual_price
                has_changed = True

            if price_item.unit_price != unit_price:
                history = PriceHistory(
                    price_item=price_item,
                    old_price=old_price,
                    new_price=new_price,
                    changed_by=changed_by or "import",
                    change_reason="Импорт XLSX",
                )
                session.add(history)
                has_changed = True

            if has_changed:
                updated += 1
            else:
                skipped += 1

    await session.commit()
    return {"created": created, "updated": updated, "skipped": skipped}
