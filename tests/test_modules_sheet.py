"""Тесты листа «🧩 Модули (распознано)» и шапки «🧩 Мебель:» (template_filler)."""
from types import SimpleNamespace

from openpyxl import Workbook

from app.services.template_filler import (
    _add_furniture_header,
    _add_modules_sheet,
    _module_ru_name,
)

SHEET = "🧩 Модули (распознано)"


def _module(type_, width, depth, height, qty=1, glass=False, facades=None, drawers=None):
    return SimpleNamespace(
        type=type_, width=width, depth=depth, height=height, quantity=qty,
        has_glass=glass, is_corner=False, facades=facades, drawers=drawers,
    )


def _room(name, modules, materials=None):
    return SimpleNamespace(room_name=name, modules=modules, materials=materials or [])


def test_module_ru_name_known_and_unknown():
    assert _module_ru_name("lower_base") == "Нижняя база"
    assert _module_ru_name("penal") == "Пенал"
    assert _module_ru_name("wardrobe") == "Шкаф-купе"
    assert _module_ru_name("unknown_xyz") == "unknown_xyz"
    assert _module_ru_name(None) == "?"


def test_modules_sheet_created_with_rows():
    wb = Workbook()
    rooms = [
        _room("Кухня", [
            _module("lower_base", 600, 560, 820, qty=2, facades={"count": 2, "type": "doors"}),
            _module("upper_base", 600, 320, 720, glass=True, facades={"count": 1, "type": "doors"}),
        ], ["EGGER H1379"]),
        _room("Ванная", [_module("vanity", 800, 500, 820, drawers={"count": 1})]),
    ]
    _add_modules_sheet(wb, rooms)
    assert SHEET in wb.sheetnames
    ws = wb[SHEET]

    # Заголовки
    assert ws["A1"].value == "Помещение"
    assert ws["B1"].value == "Мебель (модуль)"
    assert ws["C1"].value == "Габарит Ш×Г×В, мм"

    # Строки модулей
    assert ws["A2"].value == "Кухня"
    assert ws["B2"].value == "Нижняя база"
    assert ws["C2"].value == "600×560×820"
    assert ws["D2"].value == 2          # кол-во
    assert ws["E2"].value == 2          # фасадов
    assert ws["G2"].value == ""         # без стекла
    # Верхняя база со стеклом и помещение повторно не пишется (пустая ячейка)
    assert ws["A3"].value in (None, "")
    assert ws["B3"].value == "Верхняя база"
    assert ws["G3"].value == "✅"
    # Вторая комната
    assert ws["A4"].value == "Ванная"
    assert ws["B4"].value == "Тумба под раковину"
    assert ws["F4"].value == 1          # ящик


def test_modules_sheet_empty_rooms():
    wb = Workbook()
    _add_modules_sheet(wb, [])
    assert SHEET in wb.sheetnames
    assert wb[SHEET]["A2"].value == "Нет распознанных модулей"


def test_furniture_header_lists_modules():
    wb = Workbook()
    ws = wb.active
    room = _room("Кухня", [
        _module("lower_base", 600, 560, 820, qty=2, facades={"count": 2, "type": "doors"}),
        _module("upper_base", 600, 320, 720, glass=True, facades={"count": 1, "type": "doors"}),
    ], ["EGGER H1379"])
    _add_furniture_header(ws, room)

    text = ws["A2"].value
    assert text.startswith("🧩 Мебель:")
    assert "Нижняя база 600×560×820 мм — 2 шт (2 двер.)" in text
    assert "Верхняя база 600×320×720 мм — 1 шт (1 двер., стекло)" in text
    # мерж широкой шапки и высота строки
    assert "A2:H2" in [str(r) for r in ws.merged_cells.ranges]
    assert ws.row_dimensions[2].height and ws.row_dimensions[2].height > 14


def test_furniture_header_empty_room_no_crash():
    wb = Workbook()
    ws = wb.active
    _add_furniture_header(ws, _room("Кухня", []))
    assert ws["A2"].value in (None, "")
