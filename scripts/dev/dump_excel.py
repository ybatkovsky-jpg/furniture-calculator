"""Дамп содержимого Excel-сметы: имена листов + непустые строки.

Использование:
    python scripts/dev/dump_excel.py output/Расчет_xxx.xlsx [--max-cols 12] [--sheet "Лист1"]
"""
import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook


def dump(path: Path, max_cols: int, sheet_filter: str | None):
    wb = load_workbook(path, data_only=True)
    print(f"📄 Файл: {path.name} | Листы: {wb.sheetnames}")
    for ws in wb.worksheets:
        if sheet_filter and ws.title != sheet_filter:
            continue
        print(f"\n{'=' * 90}\n📋 ЛИСТ: «{ws.title}»  (строк {ws.max_row}, колонок {ws.max_column})\n{'=' * 90}")
        for row in ws.iter_rows():
            vals = [c.value for c in row[:max_cols]]
            if all(v is None or str(v).strip() == "" for v in vals):
                continue
            cells = []
            for v in vals:
                if v is None:
                    cells.append("·")
                else:
                    s = str(v)
                    cells.append(s if len(s) <= 40 else s[:37] + "...")
            print(" | ".join(cells))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", help="путь к xlsx")
    parser.add_argument("--max-cols", type=int, default=14)
    parser.add_argument("--sheet", default=None, help="только один лист")
    args = parser.parse_args()
    dump(Path(args.xlsx), args.max_cols, args.sheet)


if __name__ == "__main__":
    main()
